"""One-time load of the legacy workbook ("Active Orders on Plywood").

Reads two things: the CNC PROGRAM LIST sheet, and the plywood columns on the
queue/filed sheets (which become the first rows of the item master, with the
on-hand row as opening balances). Programs already in TPMS are left alone
unless overwrite is asked for, so the list can be re-loaded safely.
"""
import datetime
import re

import openpyxl

from util import (is_blank, make_code, norm_machine, norm_size, norm_status,
                  norm_thickness, norm_ws, parse_number, parse_time_cell,
                  title_type)

_HEADER_WORDS = {"PRO.", "#", "PROGRAM", "NUMBER"}
_TIME_COLUMNS = (("H1", 5), ("H2", 6), ("S1", 7), ("S2", 8))


def _sheet(wb, *needles):
    for ws in wb.worksheets:
        title = ws.title.upper()
        if all(n in title for n in needles):
            return ws
    return None


def _clean(value):
    return "" if is_blank(value) else norm_ws(value)


def _int(value):
    n = parse_number(value)
    return int(n) if n is not None and float(n).is_integer() and n < 100000 else None


def read_programs(wb):
    """Returns (programs, report). Each program is a dict ready for load_programs."""
    ws = _sheet(wb, "PROGRAM", "LIST")
    report = {"sheet": ws.title if ws else None, "rows": 0, "repeats": 0,
              "odd_numbers": [], "unreadable_times": 0}
    if ws is None:
        return [], report
    programs, seen = [], set()
    for idx, raw in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        r = list(raw) + [None] * (22 - len(raw))
        number = norm_ws(r[0]).upper()
        name = norm_ws(r[1])
        if not number or number in _HEADER_WORDS or name.upper() in ("NAME", "PROGRAM NAME"):
            continue
        key = (number, name)
        if key in seen:
            report["repeats"] += 1
            continue
        seen.add(key)
        flag = not re.fullmatch(r"P\d{4}", number)
        if flag:
            report["odd_numbers"].append(f"{number} (row {idx})")
        times = {}
        for code, col in _TIME_COLUMNS:
            cell = r[col]
            secs = parse_time_cell(cell)
            if secs:
                text = cell if isinstance(cell, str) else ""
                times[code] = (secs, norm_ws(text))
            elif isinstance(cell, str) and not is_blank(cell):
                report["unreadable_times"] += 1
        sheets_per_unit = r[2] if isinstance(r[2], (int, float)) and not isinstance(r[2], bool) and r[2] > 0 else None
        multi = norm_ws(r[12]).upper()
        notes = " ".join(x for x in (_clean(r[21]),) if x)
        date = r[19].date() if isinstance(r[19], datetime.datetime) and 1990 <= r[19].year <= 2100 else None
        programs.append({
            "number": number, "name": name, "description": _clean(r[20]),
            "machine_code": norm_machine(r[3]),
            "units_per_sheet": parse_number(r[4]), "units_text": _clean(r[4]),
            "sheets_per_unit": sheets_per_unit,
            "parts_on_sheet": _int(r[10]), "parts_in_unit": _int(r[11]),
            "multi_sheet": multi if multi.startswith("X") or "SHEET" in multi else "",
            "addl_machine": norm_ws(r[13]).upper() if norm_ws(r[13]).upper() in ("Y", "N") else "",
            "assembly": norm_ws(r[14]).upper() if norm_ws(r[14]).upper() in ("Y", "N") else "",
            "sheet_size": norm_size(r[15]), "thickness": norm_thickness(r[16]),
            "material_text": _clean(r[17]).upper().replace("R.BIRCH", "R. BIRCH"),
            "status": norm_status(r[18]), "program_date": date, "notes": notes,
            "number_flag": flag, "source_row": idx, "times": times,
        })
    report["rows"] = len(programs)
    return programs, report


_PROGRAM_FIELDS = ("number", "name", "description", "machine_code", "units_per_sheet", "units_text",
                   "sheets_per_unit", "parts_on_sheet", "parts_in_unit", "multi_sheet", "addl_machine",
                   "assembly", "sheet_size", "thickness", "material_text", "status", "program_date",
                   "notes", "number_flag", "source_row")


def load_programs(conn, programs, overwrite=False):
    """Insert new programs; optionally refresh existing ones. Returns counts."""
    machines = {r["code"]: r["id"] for r in conn.execute("select id, code from machine")}
    existing = {(r["number"], r["name"]): r["id"] for r in conn.execute("select id, number, name from program")}
    added = updated = skipped = times = 0
    cols = ", ".join(_PROGRAM_FIELDS)
    marks = ", ".join(["%s"] * len(_PROGRAM_FIELDS))
    sets = ", ".join(f"{f} = %s" for f in _PROGRAM_FIELDS[2:])
    for p in programs:
        values = [p[f] for f in _PROGRAM_FIELDS]
        pid = existing.get((p["number"], p["name"]))
        if pid is None:
            pid = conn.execute(f"insert into program ({cols}) values ({marks}) returning id", values).fetchone()["id"]
            added += 1
        elif overwrite:
            conn.execute(f"update program set {sets}, updated_at = now() where id = %s", values[2:] + [pid])
            conn.execute("delete from program_time where program_id = %s", (pid,))
            updated += 1
        else:
            skipped += 1
            continue
        for code, (secs, text) in p["times"].items():
            if code in machines:
                conn.execute(
                    "insert into program_time (program_id, machine_id, seconds_per_sheet, source_text) "
                    "values (%s, %s, %s, %s)", (pid, machines[code], secs, text))
                times += 1
    return {"added": added, "updated": updated, "skipped": skipped, "times": times}


def parse_item_header(header):
    """'4X8 12 MM. TXTUR GRADE QTY' -> ('4x8', '12 mm', 'Txtur'). None if not a plywood column."""
    h = norm_ws(header).upper()
    h = norm_ws(re.sub(r"\bQTY\b", "", h))
    m = re.match(r'^(\d+\s*X\s*\d+)\s+(.+?(?:MM\.?|"))\s*(.*)$', h)
    if not m:
        return None
    type_name = norm_ws(re.sub(r"\bGRADE\b", "", m.group(3)))
    type_name = re.sub(r"\bPLY$", "PLYWOOD", type_name) or "PLYWOOD"
    return norm_size(m.group(1)), norm_thickness(m.group(2)), title_type(type_name)


def read_items(wb):
    """Plywood columns from the queue and filed sheets. On-hand comes from the queue sheet's row 2."""
    found = {}
    for needles, has_on_hand in ((("CUTTING", "QUEUE"), True), (("FILED",), False)):
        ws = _sheet(wb, *needles)
        if ws is None:
            continue
        for col in range(1, ws.max_column + 1):
            parsed = parse_item_header(ws.cell(1, col).value)
            if not parsed:
                continue
            on_hand = ws.cell(2, col).value if has_on_hand else None
            on_hand = int(on_hand) if isinstance(on_hand, (int, float)) and not isinstance(on_hand, bool) else 0
            entry = found.setdefault(parsed, {"sheet_size": parsed[0], "thickness": parsed[1],
                                              "type_name": parsed[2], "on_hand": 0})
            if has_on_hand:
                entry["on_hand"] = on_hand
    return list(found.values())


def load_items(conn, items, entered_by=""):
    """Create missing types and items; give each new item an opening balance. Existing items are untouched."""
    added = balances = 0
    for it in items:
        row = conn.execute("select id from plywood_type where lower(name) = lower(%s)", (it["type_name"],)).fetchone()
        type_id = row["id"] if row else conn.execute(
            "insert into plywood_type (name) values (%s) returning id", (it["type_name"],)).fetchone()["id"]
        if conn.execute("select 1 from item where type_id = %s and thickness = %s and sheet_size = %s and grade = ''",
                        (type_id, it["thickness"], it["sheet_size"])).fetchone():
            continue
        code = base = make_code(it["sheet_size"], it["thickness"], it["type_name"])
        n = 1
        while conn.execute("select 1 from item where code = %s", (code,)).fetchone():
            n += 1
            code = f"{base}-{n}"
        item_id = conn.execute(
            "insert into item (code, type_id, thickness, sheet_size) values (%s, %s, %s, %s) returning id",
            (code, type_id, it["thickness"], it["sheet_size"])).fetchone()["id"]
        added += 1
        if it["on_hand"]:
            conn.execute(
                "insert into inventory_txn (item_id, txn_type, qty_sheets, reason, note, entered_by) "
                "values (%s, 'opening', %s, 'Opening balance', 'On-hand count from the plywood workbook', %s)",
                (item_id, it["on_hand"], entered_by))
            balances += 1
    return {"added": added, "balances": balances, "found": len(items)}


def open_workbook(file_obj):
    return openpyxl.load_workbook(file_obj, data_only=True)
